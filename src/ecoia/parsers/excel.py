from typing import Any, Dict, Optional
import openpyxl
from ecoia.parsers.base import BaseParser


class ExcelParser(BaseParser):
    SAMPLE_SIZE = 5

    def __init__(self):
        self._workbook_cache: Dict[str, Any] = {}
        self._rows_cache: Dict[str, list] = {}  # key: "filepath::sheetname"

    def _get_workbook(self, file_path: str):
        """Return cached workbook or load it once (read_only for speed)."""
        if file_path not in self._workbook_cache:
            self._workbook_cache[file_path] = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
        return self._workbook_cache[file_path]

    def _get_cached_rows(self, file_path: str, sheet_name: str) -> list:
        """Return cached rows for a sheet, loading them once from the workbook."""
        cache_key = f"{file_path}::{sheet_name}"
        if cache_key not in self._rows_cache:
            workbook = self._get_workbook(file_path)
            if sheet_name not in workbook.sheetnames:
                self._rows_cache[cache_key] = []
            else:
                sheet = workbook[sheet_name]
                self._rows_cache[cache_key] = [tuple(row) for row in sheet.iter_rows(values_only=True)]
        return self._rows_cache[cache_key]

    def close(self):
        """Close all cached workbooks and clear caches."""
        for wb in self._workbook_cache.values():
            try:
                wb.close()
            except Exception:
                pass
        self._workbook_cache.clear()
        self._rows_cache.clear()

    def parse(self, file_path: str) -> Dict[str, Any]:
        self.validate_extension(file_path, [".xlsx", ".xls"])

        workbook = self._get_workbook(file_path)
        sheets_info = {}

        for sheet_name in workbook.sheetnames:
            all_rows = self._get_cached_rows(file_path, sheet_name)
            if not all_rows:
                continue

            columns = all_rows[0]
            sample_rows = all_rows[1 : 1 + self.SAMPLE_SIZE]

            sheets_info[sheet_name] = {
                "columns": columns,
                "samples": sample_rows,
                "total_rows": len(all_rows),
            }

        return {
            "type": "excel",
            "sheets": sheets_info,
            "metadata": {"sheet_names": workbook.sheetnames, "file_path": file_path},
        }

    def get_all_rows(self, file_path: str, sheet_name: str, max_rows: Optional[int] = None):
        all_rows = self._get_cached_rows(file_path, sheet_name)
        if max_rows is not None:
            return all_rows[: max_rows + 1]
        return all_rows
