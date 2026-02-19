from typing import Any, Dict
import docx
from ecoia.parsers.base import BaseParser


class WordParser(BaseParser):
    def parse(self, file_path: str) -> Dict[str, Any]:
        self.validate_extension(file_path, [".docx"])

        doc = docx.Document(file_path)
        full_text = []
        for para in doc.paragraphs:
            full_text.append(para.text)

        return {
            "type": "word",
            "content": "\n".join(full_text),
            "metadata": {"num_paragraphs": len(doc.paragraphs)},
        }
